"""
Script de conversión de BBDD ONSV (Excel) a SQLite3 y Dataset Web Optimizado
RoadSafetyScan - SPEC-14
"""

import openpyxl
import os
import sqlite3
import json
import re
from collections import defaultdict

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_DIR = os.path.join(BASE_DIR, 'db')
OUTPUT_SQLITE = os.path.join(DB_DIR, 'roadsafety.db')
OUTPUT_WEB_JSON = os.path.join(DB_DIR, 'crashes_web.json')
OUTPUT_PUBLIC_JSON = os.path.join(BASE_DIR, 'public', 'data', 'crashes_web.json')

def sanitize_col_name(name):
    # Reemplazar acentos y caracteres especiales
    replacements = (
        ('á', 'a'), ('é', 'e'), ('í', 'i'), ('ó', 'o'), ('ú', 'u'),
        ('Á', 'a'), ('É', 'e'), ('Í', 'i'), ('Ó', 'o'), ('Ú', 'u'),
        ('ñ', 'n'), ('Ñ', 'n'), ('º', ''), ('Nº', 'num'), ('¿', ''), ('?', '')
    )
    s = str(name).strip()
    for orig, rep in replacements:
        s = s.replace(orig, rep)
    s = re.sub(r'[^\w\s]', '', s)
    s = re.sub(r'\s+', '_', s)
    return s.lower()

def read_excel_sheet(filepath, max_col_index=None):
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = None
    rows = []
    
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if idx == 5:
            # Header row
            raw_h = row if max_col_index is None else row[:max_col_index]
            header = [str(c).strip() if c is not None else f'col_{i}' for i, c in enumerate(raw_h)]
            while header and header[-1] == '':
                header.pop()
        elif idx > 5:
            if any(row):
                r = row if max_col_index is None else row[:max_col_index]
                # Slice to len(header)
                row_vals = tuple(r[i] if i < len(r) else None for i in range(len(header)))
                rows.append(row_vals)
    wb.close()
    return header, rows

def create_sqlite_database():
    print("=== Iniciando conversión a SQLite3 ===")
    if os.path.exists(OUTPUT_SQLITE):
        os.remove(OUTPUT_SQLITE)

    conn = sqlite3.connect(OUTPUT_SQLITE)
    cur = conn.cursor()

    # 1. SINIESTROS FATALES (header en fila 5, hasta col AA = 27 columnas)
    sin_file = os.path.join(DB_DIR, 'BBDD ONSV - SINIESTROS FATALES 2021-2025 (preliminar).xlsx')
    print(f"Leyendo {sin_file}...")
    h_sin, rows_sin = read_excel_sheet(sin_file, max_col_index=27)
    cols_sin = [sanitize_col_name(c) for c in h_sin]
    
    # Ensure unique column names
    seen = {}
    uniq_sin = []
    for c in cols_sin:
        if c in seen:
            seen[c] += 1
            uniq_sin.append(f"{c}_{seen[c]}")
        else:
            seen[c] = 0
            uniq_sin.append(c)

    cols_def_sin = ', '.join([f'"{c}" TEXT' for c in uniq_sin])
    cur.execute(f'CREATE TABLE siniestros ({cols_def_sin});')
    placeholders_sin = ', '.join(['?'] * len(uniq_sin))
    cur.executemany(f'INSERT INTO siniestros VALUES ({placeholders_sin})', rows_sin)
    print(f"-> Insertados {len(rows_sin)} siniestros en tabla 'siniestros'.")

    # 2. VEHICULOS
    veh_file = os.path.join(DB_DIR, 'BBDD ONSV - VEHICULOS 2021-2025 (preliminar).xlsx')
    print(f"Leyendo {veh_file}...")
    h_veh, rows_veh = read_excel_sheet(veh_file)
    cols_veh = [sanitize_col_name(c) for c in h_veh]
    seen = {}
    uniq_veh = []
    for c in cols_veh:
        if c in seen:
            seen[c] += 1
            uniq_veh.append(f"{c}_{seen[c]}")
        else:
            seen[c] = 0
            uniq_veh.append(c)

    cols_def_veh = ', '.join([f'"{c}" TEXT' for c in uniq_veh])
    cur.execute(f'CREATE TABLE vehiculos ({cols_def_veh});')
    placeholders_veh = ', '.join(['?'] * len(uniq_veh))
    cur.executemany(f'INSERT INTO vehiculos VALUES ({placeholders_veh})', rows_veh)
    print(f"-> Insertados {len(rows_veh)} vehiculos en tabla 'vehiculos'.")

    # 3. PERSONAS
    per_file = os.path.join(DB_DIR, 'BBDD ONSV - PERSONAS 2021-2025 (preliminar).xlsx')
    print(f"Leyendo {per_file}...")
    h_per, rows_per = read_excel_sheet(per_file)
    cols_per = [sanitize_col_name(c) for c in h_per]
    seen = {}
    uniq_per = []
    for c in cols_per:
        if c in seen:
            seen[c] += 1
            uniq_per.append(f"{c}_{seen[c]}")
        else:
            seen[c] = 0
            uniq_per.append(c)

    cols_def_per = ', '.join([f'"{c}" TEXT' for c in uniq_per])
    cur.execute(f'CREATE TABLE personas ({cols_def_per});')
    placeholders_per = ', '.join(['?'] * len(uniq_per))
    cur.executemany(f'INSERT INTO personas VALUES ({placeholders_per})', rows_per)
    print(f"-> Insertadas {len(rows_per)} personas en tabla 'personas'.")

    # Índices para optimizar joins
    cur.execute('CREATE INDEX idx_sin_cod ON siniestros(codigo_siniestro);')
    cur.execute('CREATE INDEX idx_veh_sin_cod ON vehiculos(codigo_siniestro);')
    cur.execute('CREATE INDEX idx_per_sin_cod ON personas(codigo_siniestro);')
    cur.execute('CREATE INDEX idx_per_gravedad ON personas(gravedad);')
    conn.commit()

    # VACUUM
    cur.execute('VACUUM;')
    conn.commit()
    conn.close()

    size_mb = os.path.getsize(OUTPUT_SQLITE) / (1024 * 1024)
    print(f"=== Base de datos SQLite creada exitosamente: {OUTPUT_SQLITE} ({size_mb:.2f} MB) ===")

def generate_web_dataset():
    print("=== Generando Dataset Web Optimizado desde SQLite ===")
    conn = sqlite3.connect(OUTPUT_SQLITE)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    # Contar vehiculos por siniestro
    cur.execute("""
        SELECT codigo_siniestro, COUNT(*) as veh_count
        FROM vehiculos
        GROUP BY codigo_siniestro
    """)
    veh_counts = {row['codigo_siniestro']: row['veh_count'] for row in cur.fetchall()}

    # Contar personas por gravedad por siniestro
    # Categorias: FALLECIDO, LESIONADO, NO SE CONOCE, ILESO
    cur.execute("""
        SELECT 
            codigo_siniestro,
            gravedad,
            COUNT(*) as count
        FROM personas
        GROUP BY codigo_siniestro, gravedad
    """)
    person_counts = defaultdict(lambda: {'fallecido': 0, 'lesionado': 0, 'noSeConoce': 0, 'ileso': 0, 'total': 0})
    for row in cur.fetchall():
        cod = row['codigo_siniestro']
        grav = (row['gravedad'] or '').strip().upper()
        cnt = row['count']
        person_counts[cod]['total'] += cnt
        if 'FALLEC' in grav:
            person_counts[cod]['fallecido'] += cnt
        elif 'LESION' in grav:
            person_counts[cod]['lesionado'] += cnt
        elif 'ILESO' in grav:
            person_counts[cod]['ileso'] += cnt
        else:
            # NO SE CONOCE u otros
            person_counts[cod]['noSeConoce'] += cnt

    # Obtener siniestros
    cur.execute("""
        SELECT 
            codigo_siniestro,
            fecha_siniestro,
            hora_siniestro,
            clase_siniestro,
            departamento,
            provincia,
            distrito,
            cod_carretera,
            red_vial,
            coordenadas_latitud,
            coordenadas_longitud
        FROM siniestros
    """)

    web_records = []
    skipped = 0

    for row in cur.fetchall():
        lat_raw = row['coordenadas_latitud']
        lon_raw = row['coordenadas_longitud']
        if lat_raw is None or lon_raw is None:
            skipped += 1
            continue
        try:
            lat = float(str(lat_raw).strip())
            lon = float(str(lon_raw).strip())
        except ValueError:
            skipped += 1
            continue

        # Validar coordenadas razonables de Perú: Lat [-19, 0], Lon [-82, -68]
        if not (-20.0 <= lat <= 1.0 and -85.0 <= lon <= -65.0):
            # Algunos archivos tienen invertidas las coordenadas o valores inválidos
            if -20.0 <= lon <= 1.0 and -85.0 <= lat <= -65.0:
                lat, lon = lon, lat
            else:
                skipped += 1
                continue

        cod = row['codigo_siniestro']
        fecha = str(row['fecha_siniestro'] or '')
        # Extraer año
        year = None
        if len(fecha) >= 4:
            # Formatos comunes: DD/MM/YYYY o YYYY-MM-DD
            m = re.search(r'\b(202[0-9])\b', fecha)
            if m:
                year = int(m.group(1))

        hora = str(row['hora_siniestro'] or '')
        p_info = person_counts.get(cod, {'fallecido': 0, 'lesionado': 0, 'noSeConoce': 0, 'ileso': 0, 'total': 0})
        v_count = veh_counts.get(cod, 0)

        # Mapear clase_siniestro a CrashType estándar
        clase = (row['clase_siniestro'] or '').strip().upper()
        crash_type = 'OTRO'
        if 'CHOQUE' in clase: crash_type = 'CHOQUE'
        elif 'ATROPELLO' in clase: crash_type = 'ATROPELLO'
        elif 'VOLCADURA' in clase: crash_type = 'VOLCADURA'
        elif 'DESPISTE' in clase: crash_type = 'DESPISTE'
        elif 'CAIDA' in clase: crash_type = 'CAIDA_PASAJERO'
        elif 'INCENDIO' in clase: crash_type = 'INCENDIO'

        # Determinar severidad principal según víctimas
        severity = 'SOLO_DANOS'
        if p_info['fallecido'] > 0:
            severity = 'FATAL'
        elif p_info['lesionado'] > 0:
            severity = 'LESIONES_GRAVES'
        elif p_info['ileso'] > 0:
            severity = 'SOLO_DANOS'

        rec = {
            'id': cod,
            'coordinates': [round(lon, 6), round(lat, 6)],
            'year': year,
            'fecha': fecha,
            'hora': hora,
            'type': crash_type,
            'clase': clase,
            'severity': severity,
            'departamento': row['departamento'],
            'provincia': row['provincia'],
            'distrito': row['distrito'],
            'roadCode': row['cod_carretera'],
            'vehiclesCount': v_count,
            'personsCount': p_info['total'],
            'personsBreakdown': {
                'fallecido': p_info['fallecido'],
                'lesionado': p_info['lesionado'],
                'noSeConoce': p_info['noSeConoce'],
                'ileso': p_info['ileso']
            }
        }
        web_records.append(rec)

    conn.close()

    print(f"Total registros para web: {len(web_records)} (Descartados sin coordenadas válidas: {skipped})")

    # Guardar en db/crashes_web.json
    with open(OUTPUT_WEB_JSON, 'w', encoding='utf-8') as f:
        json.dump(web_records, f, ensure_ascii=False, separators=(',', ':'))

    size_json_kb = os.path.getsize(OUTPUT_WEB_JSON) / 1024
    print(f"Guardado {OUTPUT_WEB_JSON} ({size_json_kb:.1f} KB)")

    # Guardar también en public/data/crashes_web.json para fetch web directo
    os.makedirs(os.path.dirname(OUTPUT_PUBLIC_JSON), exist_ok=True)
    with open(OUTPUT_PUBLIC_JSON, 'w', encoding='utf-8') as f:
        json.dump(web_records, f, ensure_ascii=False, separators=(',', ':'))
    print(f"Guardado {OUTPUT_PUBLIC_JSON} ({size_json_kb:.1f} KB)")

if __name__ == '__main__':
    create_sqlite_database()
    generate_web_dataset()
